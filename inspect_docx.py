import docx
from docx.document import Document

doc = docx.Document("Catalogo_Danone.docx")
lines = []
for para in doc.paragraphs:
    if para.text.strip():
        lines.append(para.text.strip())

for table in doc.tables:
    for row in table.rows:
        row_data = [cell.text.strip().replace("\n", " ") for cell in row.cells if cell.text.strip()]
        if row_data:
            lines.append(" | ".join(row_data))

print("Primeiras 50 linhas do documento:")
for line in lines[:50]:
    print(line)
